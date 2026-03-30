import os
import re
import json
import gradio as gr
from openai import OpenAI
from dotenv import load_dotenv
from docx import Document
from PyPDF2 import PdfReader
from fpdf import FPDF
import tempfile

# ========== 初始化 ==========
load_dotenv()

client = OpenAI(
    api_key=os.getenv("OPENAI_API_KEY"),
    base_url=os.getenv("OPENAI_BASE_URL"),
)
MODEL = os.getenv("OPENAI_MODEL", "deepseek-chat")

# 尝试加载中文字体（用于PDF导出）
FONT_PATH = None
for p in [
    "C:/Windows/Fonts/msyh.ttc",
    "C:/Windows/Fonts/simsun.ttc",
    "C:/Windows/Fonts/simhei.ttf",
]:
    if os.path.exists(p):
        FONT_PATH = p
        break


# ========== 文件读取 ==========
def read_file(file) -> str:
    if file is None:
        return ""
    filepath = file.name if hasattr(file, "name") else file
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext == ".txt":
            with open(filepath, "r", encoding="utf-8") as f:
                return f.read()
        elif ext == ".docx":
            doc = Document(filepath)
            return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        elif ext == ".pdf":
            reader = PdfReader(filepath)
            return "\n".join([page.extract_text() or "" for page in reader.pages])
        else:
            return ""
    except Exception:
        return ""


# ========== AI 调用 ==========
def call_ai(system: str, user: str) -> str:
    response = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
        temperature=0.7,
    )
    return response.choices[0].message.content


# ========== 功能1: 简历评分 ==========
def get_score(resume: str, jd: str) -> dict:
    """返回评分和各项维度得分"""
    prompt = f"""请对这份简历与目标岗位的匹配度进行评分。

目标岗位JD：
{jd}

简历内容：
{resume}

请严格按以下JSON格式返回，不要添加任何其他内容：
{{
    "total_score": 85,
    "dimensions": {{
        "技能匹配": 80,
        "经验匹配": 75,
        "教育背景": 90,
        "表达质量": 85
    }},
    "summary": "一句话总评"
}}"""
    try:
        result = call_ai("你是一位专业的简历评估专家。只返回JSON，不要多余内容。", prompt)
        # 提取JSON
        json_match = re.search(r'\{[\s\S]*\}', result)
        if json_match:
            return json.loads(json_match.group())
    except Exception:
        pass
    return {"total_score": 0, "dimensions": {}, "summary": "评分失败"}


# ========== 功能2: ATS关键词匹配 ==========
def get_keywords(resume: str, jd: str) -> dict:
    """分析关键词匹配情况"""
    prompt = f"""分析简历与岗位JD的关键词匹配情况。

目标岗位JD：
{jd}

简历内容：
{resume}

请严格按以下JSON格式返回，不要添加任何其他内容：
{{
    "matched_keywords": ["Python", "机器学习"],
    "missing_keywords": ["TensorFlow", "A/B测试"],
    "coverage_rate": 65,
    "suggestions": "建议补充以下关键词..."
}}"""
    try:
        result = call_ai("你是一位ATS简历系统专家。只返回JSON，不要多余内容。", prompt)
        json_match = re.search(r'\{[\s\S]*\}', result)
        if json_match:
            return json.loads(json_match.group())
    except Exception:
        pass
    return {"matched_keywords": [], "missing_keywords": [], "coverage_rate": 0, "suggestions": ""}


# ========== 功能3: 详细优化建议 ==========
def get_detailed_analysis(resume: str, jd: str) -> str:
    """生成详细分析和优化建议"""
    prompt = f"""你是一位资深HR和简历优化专家。请根据目标岗位要求，详细分析这份简历。

目标岗位JD：
{jd}

当前简历：
{resume}

请按以下结构输出（使用Markdown格式）：

## 📊 简历优缺点分析
（从岗位匹配度、技能覆盖、经历描述、教育背景等维度分析）

## ✏️ 针对性修改建议
（给出具体的、可执行的修改建议，每条建议说明改什么、怎么改）

## 📄 优化后的简历
（直接给出修改后的完整简历内容）

## 🎤 面试准备建议
（根据简历和JD，预测可能的面试问题，并给出回答建议）

请用中文回答。"""
    return call_ai("你是一位专业的简历优化顾问，回答要专业、具体、有针对性。", prompt)


# ========== 功能4: 多岗位对比 ==========
def compare_jobs(resume: str, jd_list: list) -> str:
    """一份简历对比多个岗位"""
    jd_text = "\n\n".join([f"### 岗位{i+1}：{jd}" for i, jd in enumerate(jd_list) if jd.strip()])
    if not jd_text.strip():
        return "请至少输入一个岗位JD"

    prompt = f"""一份简历对比以下多个目标岗位，分析每个岗位的匹配度。

简历内容：
{resume}

目标岗位列表：
{jd_text}

请对每个岗位给出：
1. 匹配度评分（满分100）
2. 优势分析
3. 短板分析
4. 建议（是否推荐投递）

最后给出总体建议：最适合投递哪个岗位。

请用中文回答，使用Markdown格式。"""
    return call_ai("你是一位资深HR，擅长岗位匹配分析。", prompt)


# ========== 功能5: PDF导出 ==========
def export_pdf(content: str) -> str:
    """将分析结果导出为PDF"""
    try:
        pdf = FPDF()
        pdf.add_page()

        if FONT_PATH:
            pdf.add_font("Chinese", "", FONT_PATH, uni=True)
            pdf.set_font("Chinese", size=11)
        else:
            pdf.set_font("Helvetica", size=11)

        pdf.set_auto_page_break(auto=True, margin=15)

        # 标题
        pdf.set_font_size(16)
        if FONT_PATH:
            pdf.set_font("Chinese", size=16)
        pdf.cell(0, 15, "AI 简历优化报告", new_x="LMARGIN", new_y="NEXT", align="C")
        pdf.ln(5)

        # 内容
        pdf.set_font_size(11)
        if FONT_PATH:
            pdf.set_font("Chinese", size=11)

        # 清理Markdown符号，保留纯文本
        clean_content = re.sub(r'[#*`>|]', '', content)
        for line in clean_content.split('\n'):
            line = line.strip()
            if line:
                pdf.multi_cell(0, 7, line)
                pdf.ln(2)

        # 保存到临时文件
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
        pdf.output(tmp.name)
        return tmp.name
    except Exception as e:
        return ""


# ========== 主处理函数 ==========
def process_all(resume_text, resume_file, jd1, jd2, jd3):
    """一键处理所有功能"""

    # 获取简历内容
    if resume_file is not None:
        resume = read_file(resume_file)
        if not resume.strip():
            yield "⚠️ 无法读取上传的文件", "", "", "", None, gr.update(visible=False)
            return
    else:
        resume = resume_text

    if not resume.strip():
        yield "⚠️ 请先输入或上传你的简历", "", "", "", None, gr.update(visible=False)
        return
    if not jd1.strip():
        yield "⚠️ 请先至少输入一个目标岗位JD", "", "", "", None, gr.update(visible=False)
        return

    yield "⏳ 正在分析中，请稍候...", "", "", "", None, gr.update(visible=False)

    # 1. 简历评分
    score_data = get_score(resume, jd1)
    score_md = format_score(score_data)
    yield score_md, "⏳ 分析中...", "", "", None, gr.update(visible=False)

    # 2. ATS关键词
    kw_data = get_keywords(resume, jd1)
    kw_md = format_keywords(kw_data)
    yield score_md, kw_md, "⏳ 分析中...", "", None, gr.update(visible=False)

    # 3. 详细分析
    analysis = get_detailed_analysis(resume, jd1)
    yield score_md, kw_md, analysis, "⏳ 分析中...", None, gr.update(visible=False)

    # 4. 多岗位对比（如果有多个JD）
    jd_list = [jd for jd in [jd1, jd2, jd3] if jd.strip()]
    if len(jd_list) > 1:
        comparison = compare_jobs(resume, jd_list)
    else:
        comparison = "> 仅输入了一个岗位，跳过多岗位对比。\n> 输入多个岗位JD可进行对比分析。"
    yield score_md, kw_md, analysis, comparison, None, gr.update(visible=False)

    # 5. PDF导出
    pdf_content = f"# 简历评分\n{score_md}\n\n# ATS关键词分析\n{kw_md}\n\n# 详细分析\n{analysis}\n\n# 多岗位对比\n{comparison}"
    pdf_path = export_pdf(pdf_content)
    if pdf_path:
        yield score_md, kw_md, analysis, comparison, pdf_path, gr.update(visible=True)
    else:
        yield score_md, kw_md, analysis, comparison, None, gr.update(visible=False)


def format_score(data: dict) -> str:
    score = data.get("total_score", 0)
    summary = data.get("summary", "")
    dims = data.get("dimensions", {})

    # 分数颜色
    if score >= 80:
        emoji = "🟢"
    elif score >= 60:
        emoji = "🟡"
    else:
        emoji = "🔴"

    md = f"## {emoji} 综合评分：{score}/100\n\n"
    if summary:
        md += f"> {summary}\n\n"

    if dims:
        md += "| 维度 | 得分 |\n|------|------|\n"
        for k, v in dims.items():
            bar = "█" * (v // 10) + "░" * (10 - v // 10)
            md += f"| {k} | {bar} {v} |\n"

    return md


def format_keywords(data: dict) -> str:
    matched = data.get("matched_keywords", [])
    missing = data.get("missing_keywords", [])
    rate = data.get("coverage_rate", 0)
    suggestions = data.get("suggestions", "")

    md = f"## 🔑 ATS 关键词匹配：{rate}%\n\n"

    if matched:
        md += f"**✅ 已匹配（{len(matched)}个）：** "
        md += "  ".join([f"`{k}`" for k in matched])
        md += "\n\n"

    if missing:
        md += f"**❌ 缺失（{len(missing)}个）：** "
        md += "  ".join([f"`{k}`" for k in missing])
        md += "\n\n"

    if suggestions:
        md += f"**💡 建议：** {suggestions}\n"

    return md


# ========== 网页界面 ==========
with gr.Blocks(title="AI 简历优化器 v2.0", theme=gr.themes.Soft()) as demo:

    gr.Markdown("# 📄 AI 简历优化器 v2.0")
    gr.Markdown("五大功能：智能评分 · ATS匹配 · 优化建议 · 多岗对比 · PDF导出")

    with gr.Row():
        # 左侧：输入
        with gr.Column(scale=1):
            gr.Markdown("### 📝 你的简历")
            with gr.Tab("手动输入"):
                resume_text = gr.Textbox(
                    placeholder="粘贴简历内容...", lines=12, show_label=False)
            with gr.Tab("上传文件"):
                resume_file = gr.File(
                    label="支持 .txt / .docx / .pdf",
                    file_types=[".txt", ".docx", ".pdf"])

            gr.Markdown("### 🎯 目标岗位 JD（可填1-3个）")
            jd1 = gr.Textbox(label="岗位 1（必填）", placeholder="粘贴岗位1的JD...", lines=6)
            with gr.Accordion("➕ 添加更多岗位（可选）", open=False):
                jd2 = gr.Textbox(label="岗位 2", placeholder="粘贴岗位2的JD...", lines=5)
                jd3 = gr.Textbox(label="岗位 3", placeholder="粘贴岗位3的JD...", lines=5)

            submit_btn = gr.Button("🚀 一键分析", variant="primary", size="large")

        # 右侧：输出
        with gr.Column(scale=1):
            with gr.Tabs():
                with gr.Tab("📊 评分"):
                    score_output = gr.Markdown(value="> 等待分析...")
                with gr.Tab("🔑 关键词"):
                    keyword_output = gr.Markdown(value="> 等待分析...")
                with gr.Tab("💡 优化建议"):
                    analysis_output = gr.Markdown(value="> 等待分析...")
                with gr.Tab("⚖️ 多岗对比"):
                    comparison_output = gr.Markdown(value="> 等待分析...")

            pdf_output = gr.File(label="📥 导出报告", visible=False)

    gr.Markdown(
        "<div style='text-align:center;color:#999;margin-top:16px;'>"
        "⭐ 觉得好用？欢迎到 <a href='https://github.com/tangp333/ai-resume-optimizer' target='_blank'>GitHub</a> 点个 Star"
        "  |  Powered by AI"
        "</div>"
    )

    submit_btn.click(
        fn=process_all,
        inputs=[resume_text, resume_file, jd1, jd2, jd3],
        outputs=[score_output, keyword_output, analysis_output, comparison_output, pdf_output, pdf_output],
    )

if __name__ == "__main__":
    demo.launch(share=True)
