import os
import re
import json
import tempfile
from flask import Flask, request, jsonify, render_template, send_file
from openai import OpenAI
from dotenv import load_dotenv
from docx import Document
from PyPDF2 import PdfReader
from fpdf import FPDF
from werkzeug.utils import secure_filename

load_dotenv()

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB

client = OpenAI(
    api_key=os.getenv("OPENAI_API_KEY"),
    base_url=os.getenv("OPENAI_BASE_URL"),
)
MODEL = os.getenv("OPENAI_MODEL", "deepseek-chat")

# 中文字体
FONT_PATH = None
for p in ["C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simsun.ttc", "C:/Windows/Fonts/simhei.ttf"]:
    if os.path.exists(p):
        FONT_PATH = p
        break


def read_file(file) -> str:
    filename = secure_filename(file.filename)
    ext = os.path.splitext(filename)[1].lower()
    try:
        if ext == ".txt":
            return file.read().decode("utf-8")
        elif ext == ".docx":
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
            file.save(tmp.name)
            doc = Document(tmp.name)
            return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        elif ext == ".pdf":
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
            file.save(tmp.name)
            reader = PdfReader(tmp.name)
            return "\n".join([page.extract_text() or "" for page in reader.pages])
    except Exception as e:
        return ""
    return ""


def call_ai(system: str, user: str) -> str:
    response = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
        temperature=0.7,
    )
    return response.choices[0].message.content


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/analyze", methods=["POST"])
def analyze():
    resume_text = request.form.get("resume", "").strip()
    jd = request.form.get("jd", "").strip()

    # 处理文件上传
    if "file" in request.files and request.files["file"].filename:
        file_content = read_file(request.files["file"])
        if file_content:
            resume_text = file_content

    if not resume_text:
        return jsonify({"error": "请输入或上传简历"}), 400
    if not jd:
        return jsonify({"error": "请输入岗位JD"}), 400

    # 1. 评分
    score = get_score(resume_text, jd)
    # 2. 关键词
    keywords = get_keywords(resume_text, jd)
    # 3. 详细分析
    analysis = get_analysis(resume_text, jd)

    return jsonify({
        "score": score,
        "keywords": keywords,
        "analysis": analysis,
    })


@app.route("/api/compare", methods=["POST"])
def compare():
    data = request.get_json()
    resume = data.get("resume", "").strip()
    jds = [j.strip() for j in data.get("jds", []) if j.strip()]

    if not resume or len(jds) < 2:
        return jsonify({"error": "需要简历和至少2个岗位JD"}), 400

    jd_text = "\n\n".join([f"### 岗位{i+1}：{j}" for i, j in enumerate(jds)])
    prompt = f"""一份简历对比多个目标岗位，分析每个岗位的匹配度。

简历：{resume}

岗位列表：
{jd_text}

对每个岗位给出：匹配度评分（满分100）、优势、短板、是否推荐投递。
最后给出最适合投递哪个岗位的建议。用中文Markdown格式。"""

    result = call_ai("你是一位资深HR，擅长岗位匹配分析。", prompt)
    return jsonify({"result": result})


@app.route("/api/export-pdf", methods=["POST"])
def export_pdf():
    data = request.get_json()
    content = data.get("content", "")

    try:
        pdf = FPDF()
        pdf.add_page()
        if FONT_PATH:
            pdf.add_font("Chinese", "", FONT_PATH, uni=True)
            pdf.set_font("Chinese", size=12)
        else:
            pdf.set_font("Helvetica", size=12)

        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.set_font_size(18)
        if FONT_PATH:
            pdf.set_font("Chinese", size=18)
        pdf.cell(0, 15, "AI 简历优化报告", new_x="LMARGIN", new_y="NEXT", align="C")
        pdf.ln(8)
        pdf.set_font_size(11)
        if FONT_PATH:
            pdf.set_font("Chinese", size=11)

        clean = re.sub(r'[#*`>|]', '', content)
        for line in clean.split('\n'):
            line = line.strip()
            if line:
                pdf.multi_cell(0, 7, line)
                pdf.ln(2)

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
        pdf.output(tmp.name)
        return send_file(tmp.name, as_attachment=True, download_name="简历优化报告.pdf")
    except Exception as e:
        return jsonify({"error": str(e)}), 500


def get_score(resume, jd):
    prompt = f"""对简历与岗位的匹配度评分。

JD：{jd}
简历：{resume}

严格返回JSON：
{{"total_score":85,"dimensions":{{"技能匹配":80,"经验匹配":75,"教育背景":90,"表达质量":85}},"summary":"一句话总评"}}"""
    try:
        result = call_ai("简历评估专家。只返回JSON。", prompt)
        m = re.search(r'\{[\s\S]*\}', result)
        if m:
            return json.loads(m.group())
    except:
        pass
    return {"total_score": 0, "dimensions": {}, "summary": "评分失败"}


def get_keywords(resume, jd):
    prompt = f"""分析简历与JD的关键词匹配。

JD：{jd}
简历：{resume}

严格返回JSON：
{{"matched":["Python"],"missing":["TensorFlow"],"coverage_rate":65,"suggestions":"建议补充..."}}"""
    try:
        result = call_ai("ATS专家。只返回JSON。", prompt)
        m = re.search(r'\{[\s\S]*\}', result)
        if m:
            return json.loads(m.group())
    except:
        pass
    return {"matched": [], "missing": [], "coverage_rate": 0, "suggestions": ""}


def get_analysis(resume, jd):
    prompt = f"""资深HR分析简历。JD：{jd}\n简历：{resume}\n\n输出：\n## 📊 优缺点分析\n## ✏️ 修改建议\n## 📄 优化后简历\n## 🎤 面试准备\n\n中文Markdown。"""
    return call_ai("简历优化顾问，专业具体。", prompt)


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=7860)
